from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'/'submission'/'StudentBridge_AI_Presentation.pptx'
rs=Presentation();rs.slide_width=Inches(13.33);rs.slide_height=Inches(7.5)
C={'dark':'113438','navy':'0B282D','mint':'DFFF9A','green':'327B65','cream':'F7F9F3','white':'FFFFFF','gray':'80938B','line':'DDE7DF','light':'E9F4E9','orange':'BD7949'}
def rgb(hex):return RGBColor.from_string(C.get(hex,hex))
def add_text(slide,x,y,w,h,text,size=16,bold=False,color='dark',font='Aptos',align=None):
 tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True
 tf.margin_left=Inches(.03);tf.margin_top=Inches(.02);tf.margin_right=Inches(.03);tf.margin_bottom=Inches(.01)
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.space_after=Pt(3)
  p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
  if align:p.alignment=align
 return tb

def shape(slide,x,y,w,h,fill='white',stroke=None,r=0.15):
 t=MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE
 obj=slide.shapes.add_shape(t,Inches(x),Inches(y),Inches(w),Inches(h));obj.fill.solid();obj.fill.fore_color.rgb=rgb(fill);obj.line.fill.background() if not stroke else None
 if stroke:obj.line.color.rgb=rgb(stroke);obj.line.width=Pt(1)
 return obj

def base(n,title,kicker=None):
 slide=rs.slides.add_slide(rs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb('cream')
 shape(slide,0,0,13.33,.095,fill='dark',r=0)
 add_text(slide,.68,.26,6,.27,'STUDENTBRIDGE AI  /  TOPIC C',9,True,'green')
 add_text(slide,.68,.75,12,1,title,30,True)
 if kicker:add_text(slide,.71,1.65,11.7,.55,kicker,13,False,'gray')
 shape(slide,.68,7.18,11.97,.015,'line',r=0)
 add_text(slide,.68,7.2,8,.19,'2026 SHENZHEN INTERNATIONAL FINTECH COMPETITION',8,True,'gray')
 add_text(slide,12.15,7.17,.5,.20,f'{n:02d}',9,True,'green',align=PP_ALIGN.RIGHT)
 return slide

def tile(slide,x,y,w,h,num,title,desc,accent='light'):
 shape(slide,x,y,w,h,'white','line')
 shape(slide,x+.19,y+.18,.47,.47,'light')
 add_text(slide,x+.27,y+.25,.4,.3,str(num),11,True,'green')
 add_text(slide,x+.2,y+.84,w-.4,.42,title,16,True,'dark')
 add_text(slide,x+.2,y+1.38,w-.4,h-1.55,desc,11,False,'gray')

# cover
sl=rs.slides.add_slide(rs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb('dark')
shape(sl,8.2,0,5.13,7.5,'navy')
for i in range(3):
 o=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(7.65+i*.57),Inches(.9+i*.47),Inches(4.4-i*.25),Inches(4.4-i*.25));o.fill.background();o.line.color.rgb=rgb('green');o.line.width=Pt(.9)
shape(sl,.72,.7,1.8,.40,'mint');add_text(sl,.9,.79,1.4,.2,'TOPIC C',10,True,'dark')
add_text(sl,.74,2.06,7.5,1.68,'Finance without\nfrontiers.',46,True,'white')
add_text(sl,.79,4.17,6.9,.85,'An explainable student finance agent for 6 countries and 30 cross-border directions.',18,False,'mint')
add_text(sl,.80,6.62,10,.3,'STUDENTBRIDGE AI  ·  DEMONSTRATION SANDBOX  ·  OCTOBER 2026',10,True,'white')
add_text(sl,9.22,3.00,2.2,1.0,'✳',74,True,'mint',align=PP_ALIGN.CENTER)

# 2
sl=base(2,'The problem is not currency. It is complexity.','Students face five connected decisions without a single safe workflow.')
for i,(a,b) in enumerate([('01','Tuition deadlines'),('02','Currency spreads and fees'),('03','Route eligibility'),('04','Living-cost uncertainty'),('05','Building local credit')]):
 x=.76+(i%3)*4.15;y=2.22+(i//3)*1.9
 shape(sl,x,y,3.9,1.62,'white','line');add_text(sl,x+.21,y+.18,.6,.3,a,13,True,'green');add_text(sl,x+.2,y+.67,3.50,.73,b,18,True,'dark')
add_text(sl,9.13,5.92,3,.8,'One financial journey.\nMany hidden constraints.',13,True,'green')

# 3
sl=base(3,'The six-country corridor network','Any country can be a sending country OR a study destination; we do not assume every bank route exists.')
flags=[('CHINA','CNY'),('UNITED KINGDOM','GBP'),('JAPAN','JPY'),('SOUTH KOREA','KRW'),('AUSTRALIA','AUD'),('BANGLADESH','BDT')]
for i,(country,cur) in enumerate(flags):
 x=.75+(i%3)*4.15;y=2.13+(i//3)*1.88
 shape(sl,x,y,3.9,1.55,'white','line');add_text(sl,x+.23,y+.27,3.5,.37,country,18,True,'dark');add_text(sl,x+.24,y+.91,3.2,.3,cur+'   ·   SEND OR RECEIVE',11,True,'green')
add_text(sl,7.95,6.2,4.6,.43,'6 × 5 = 30 ORDERED PAIRS',17,True,'dark',align=PP_ALIGN.RIGHT)

# 4
sl=base(4,'An end-to-end financial agent workflow','Task completion is a competition priority (40%).')
steps=[('1','Set your journey','Origin, destination and programme'),('2','Read the invoice','Exact amount, deadline, funding'),('3','Check the data','Reference FX, user quotes, source links'),('4','Apply the rules','Cost, deadline and reserve constraints'),('5','Explain every choice','Decision Passport and Digital Twin'),('6','Approve the sandbox','Simulated receipt and audit event')]
for i,(no,title,desc) in enumerate(steps):
 x=.75+(i%3)*4.15;y=2.04+(i//3)*2.19
 tile(sl,x,y,3.91,1.99,no,title,desc)

# 5
sl=base(5,'Our innovation: Three connected decisions','These features work together, not as three separate chatbot answers.')
for i,(num,title,desc) in enumerate([('01','EduRoute Intelligence','Eligibility, available funds, protected reserve and tuition deadline are hard constraints. Rank only eligible options.'),('02','Financial Digital Twin','Change FX, unexpected fees or living expenses. See the new budget and whether your reserve survives.'),('03','Decision Passport','Show every calculation, source, rejected alternative, warning and a plan integrity digest.')]):
 x=.75+i*4.15;tile(sl,x,2.24,3.92,3.56,num,title,desc)
shape(sl,.75,6.15,11.9,.59,'dark');add_text(sl,1.05,6.28,11.3,.3,'Innovation = auditable financial decisions + explicit uncertainty, not unexplained AI output.',14,True,'mint',align=PP_ALIGN.CENTER)

# 6
sl=base(6,'Financial logic that can be tested','The AI does not invent a bank rate or change the decision engine arithmetic.')
shape(sl,.8,2.1,7.3,3.8,'white','line')
add_text(sl,1.13,2.36,6.7,.33,'DETERMINISTIC COST MODEL',11,True,'green')
add_text(sl,1.13,2.98,6.5,.75,'Required tuition = Invoice − confirmed scholarship',21,True,'dark')
add_text(sl,1.13,3.86,6.5,.75,'Total debit = Tuition × quoted rate + applicable fees',20,True,'dark')
add_text(sl,1.13,4.83,6.5,.8,'Then check: eligibility, deadline, funds after reserve and quote expiry.',15,False,'gray')
shape(sl,8.35,2.1,4.14,3.8,'dark')
add_text(sl,8.68,2.55,3.5,.4,'CRITICAL REFUSAL CASE',12,True,'mint')
add_text(sl,8.68,3.27,3.48,1.2,'No verified or entered suitable route?',22,True,'white')
add_text(sl,8.68,4.74,3.45,.85,'The system explains the gap. It does not fabricate a recommendation.',14,False,'mint')

# 7
sl=base(7,'Real facts. Transparent gaps.','Source provenance is a product feature, not an afterthought.')
rows=[('LIVE REFERENCE FX','Frankfurter v2; dated, reference-only, not a transfer quote.'),('PUBLISHED UNIVERSITY FEES','Official 2026 programme schedules where amount is directly supported.'),('STUDENT PERSONAL FINANCE','User-entered, self-reported, not independently verified.'),('PROVIDER TRANSFER QUOTES','User-entered and explicitly unverified without provider integration.'),('EXECUTION EVIDENCE','MySQL sandbox log and atomic audit chain; not bank settlement.')]
for i,(title,desc) in enumerate(rows):
 y=2.05+i*.89;shape(sl,.76,y,11.78,.74,'white','line');add_text(sl,1.02,y+.13,3.38,.45,title,12,True,'green');add_text(sl,4.52,y+.15,7.6,.45,desc,12,False,'dark')

# 8
sl=base(8,'Security is part of the workflow','We prioritize constraints and limitations over unsupported compliance claims.')
for i,(title,desc) in enumerate([('Permission tiers','Read data; plan; explicit sandbox approver; local operator.'),('Human approval','Approval confirmation is required for every sandbox execution.'),('Safer inputs','Pydantic validation; PDF size/type limits; user quote provenance.'),('Audit proof','HMAC-chained events, SHA-256 passports, idempotent test payments.')]):
 x=.78+(i%2)*6.2;y=2.04+(i//2)*2.1;shape(sl,x,y,5.88,1.85,'white','line');add_text(sl,x+.25,y+.33,5.39,.33,title,18,True,'dark');add_text(sl,x+.25,y+.86,5.39,.68,desc,12,False,'gray')
add_text(sl,.82,6.63,11.6,.29,'NOT production-ready: no KYC/AML certification, user authentication, or live bank transaction execution.',10,True,'orange')

# 9
sl=base(9,'MySQL-powered demonstration and test evidence','Automated tests, two example corridors, and an auditable simulated execution.')
for i,(num,name,desc) in enumerate([('32','TESTS PASSING','Deadline, fee, eligibility, expiry, idempotency, API validation.'),('02','SANDBOX CASES','China → Australia and Japan → United Kingdom.'),('0','REAL PAYMENTS','Only explicit, local, approved test ledger entries.')]):
 x=.76+i*4.12;shape(sl,x,2.19,3.88,2.30,'white','line');add_text(sl,x+.25,2.52,3.25,.7,num,42,True,'green');add_text(sl,x+.25,3.43,3.45,.31,name,12,True,'dark');add_text(sl,x+.25,3.84,3.35,.52,desc,10,False,'gray')
shape(sl,.76,5.07,12.0,1.14,'dark');add_text(sl,1.10,5.36,11.2,.58,'MySQL integration smoke test and real screen-recorded demonstration still required.',14,True,'mint')

# 10
sl=base(10,'A responsible foundation for a scalable product','Next: authenticated bank quote APIs, verified corridor compliance, and secure user accounts.')
add_text(sl,1.11,2.17,10.4,1.35,'Students deserve financial confidence,\nnot another black-box recommendation.',31,True,'dark')
shape(sl,1.10,4.53,11.06,1.28,'dark');add_text(sl,1.54,4.79,10.2,.55,'STUDENTBRIDGE AI   /   EXPLAIN THE COST. PROTECT THE STUDENT. SHOW THE EVIDENCE.',16,True,'mint',align=PP_ALIGN.CENTER)
add_text(sl,1.10,6.38,11.1,.3,'SOURCE CODE  •  TECHNICAL REPORT  •  SECURITY SELF-ASSESSMENT  •  SANDBOX LOGS',12,True,'green')

OUT.parent.mkdir(parents=True,exist_ok=True)
rs.save(OUT)
print('slides',len(rs.slides),'bytes',OUT.stat().st_size,'saved',OUT)
